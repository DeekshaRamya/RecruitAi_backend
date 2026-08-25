import io
import logging
import os
import re
from pypdf import PdfReader
import docx

logger = logging.getLogger("recruitai-backend.file_parser")

def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:    
        pdf_file = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_file)
        text = ""
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        if text.strip():
            return text.strip()
    except Exception as e:
        logger.warning(f"PyPDF extraction warning: {e}")
    
    # Fallback to readable string parsing
    try:
        raw_str = file_bytes.decode("latin-1", errors="ignore")
        words = re.findall(r'[a-zA-Z0-9.,@#+-_/]{2,}', raw_str)
        return " ".join(words[:600])
    except Exception:
        return ""

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        docx_file = io.BytesIO(file_bytes)
        doc = docx.Document(docx_file)
        paragraphs = [p.text for p in doc.paragraphs]
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_text.append(cell.text)
        full_text = "\n".join(paragraphs + table_text)
        if full_text.strip():
            return full_text.strip()
    except Exception as e:
        logger.warning(f"DOCX extraction warning: {e}")
    
    try:
        return file_bytes.decode("utf-8", errors="ignore").strip()
    except Exception:
        return ""

def extract_text(filename: str, file_bytes: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    extracted = ""
    if ext == ".pdf":
        extracted = extract_text_from_pdf(file_bytes)
    elif ext in [".docx", ".doc"]:
        extracted = extract_text_from_docx(file_bytes)
    else:
        try:
            extracted = file_bytes.decode("utf-8", errors="ignore").strip()
        except Exception:
            extracted = ""
    
    return extracted.strip()
